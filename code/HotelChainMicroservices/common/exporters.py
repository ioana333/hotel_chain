from __future__ import annotations
from abc import ABC, abstractmethod
from pathlib import Path
import csv
import json
import xml.etree.ElementTree as ET

class IExporter(ABC):
    @abstractmethod
    def export(self, rows: list[dict], output_path: str) -> str:
        raise NotImplementedError

class CSVExporter(IExporter):
    def export(self, rows: list[dict], output_path: str) -> str:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open('w', newline='', encoding='utf-8') as f:
            if not rows:
                f.write('')
                return str(path)
            writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        return str(path)

class JSONExporter(IExporter):
    def export(self, rows: list[dict], output_path: str) -> str:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding='utf-8')
        return str(path)

class XMLExporter(IExporter):
    def export(self, rows: list[dict], output_path: str) -> str:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        root = ET.Element('items')
        for row in rows:
            item = ET.SubElement(root, 'item')
            for key, value in row.items():
                child = ET.SubElement(item, key)
                child.text = '' if value is None else str(value)
        ET.ElementTree(root).write(path, encoding='utf-8', xml_declaration=True)
        return str(path)

class DOCExporter(IExporter):
    def export(self, rows: list[dict], output_path: str) -> str:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            from docx import Document
            document = Document()
            document.add_heading('Hotel Chain Export', level=1)
            if rows:
                table = document.add_table(rows=1, cols=len(rows[0]))
                for i, key in enumerate(rows[0].keys()):
                    table.rows[0].cells[i].text = key
                for row in rows:
                    cells = table.add_row().cells
                    for i, value in enumerate(row.values()):
                        cells[i].text = '' if value is None else str(value)
            else:
                document.add_paragraph('Nu există date de exportat.')
            document.save(path)
        except Exception:
            # fallback: fișier .doc textual, ca să funcționeze chiar dacă python-docx lipsește
            path.write_text('\n'.join(str(row) for row in rows), encoding='utf-8')
        return str(path)

class ExporterFactory:
    """Factory Method pentru crearea exportatorului potrivit fără if-uri în servicii."""
    @staticmethod
    def create(fmt: str) -> IExporter:
        fmt = fmt.lower().strip('.')
        exporters = {
            'csv': CSVExporter,
            'json': JSONExporter,
            'xml': XMLExporter,
            'doc': DOCExporter,
            'docx': DOCExporter,
        }
        if fmt not in exporters:
            raise ValueError('Format acceptat: csv, json, xml, doc')
        return exporters[fmt]()


def keep_fields(rows: list[dict], fields: list[str]) -> list[dict]:
    """Păstrează în export doar câmpurile cerute de funcționalitatea respectivă."""
    return [{field: row.get(field, '') for field in fields} for row in rows]
